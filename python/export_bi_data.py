#!/usr/bin/env python3
"""
export_bi_data.py - export the reporting views for Power BI / Tableau / Excel

Reads the views defined in sql/03_views.sql and writes:

  data/bi/*.csv                          one CSV per dataset
  data/bi/ServiceOps_Dashboard_Data.xlsx the same datasets as sheets + a data dictionary

Both Power BI ("Get data > Excel workbook") and Tableau Public ("Connect > Microsoft
Excel") read the .xlsx directly. See dashboard/POWER_BI_GUIDE.md and
dashboard/TABLEAU_GUIDE.md.

Usage
  pip install pymysql openpyxl
  python3 python/export_bi_data.py            # uses MYSQL_HOST / MYSQL_USER / MYSQL_PASSWORD env vars
"""

import csv
import datetime as dt
import decimal
import os
from pathlib import Path

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "bi"

DATASETS = {
    # sheet name: (query, description)
    "WorkOrders": (
        "SELECT * FROM v_work_order_summary ORDER BY work_order_id",
        "Main fact table. One row per work order (repair visit): timing, revenue, cost, profit, "
        "parts delays, comebacks and the customer's rating.",
    ),
    "WorkOrderParts": (
        """SELECT wop.work_order_id, DATE(w.opened_at) AS opened_date, p.sku, p.part_name, p.part_category,
                  s.supplier_name, wop.quantity, wop.unit_cost, wop.unit_price,
                  wop.quantity * wop.unit_price AS line_revenue, wop.quantity * wop.unit_cost AS line_cost,
                  wop.was_in_stock, (p.reorder_qty = 0) AS is_special_order_only
           FROM work_order_part wop
           JOIN work_order w ON w.work_order_id = wop.work_order_id
           JOIN part p       ON p.part_id = wop.part_id
           JOIN supplier s   ON s.supplier_id = p.supplier_id
           ORDER BY wop.work_order_id, p.sku""",
        "One row per part installed on a job. Links to WorkOrders on work_order_id.",
    ),
    "Appointments": (
        "SELECT * FROM v_appointment_summary ORDER BY appointment_id",
        "One row per appointment: booking lead time, slot, outcome (arrived / no-show / cancelled / "
        "scheduled) and the customer's wait at check-in.",
    ),
    "PartsInventory": (
        "SELECT * FROM v_part_inventory ORDER BY part_id",
        "One row per part as of 2025-12-31: stock on hand, on order, usage, stockouts and the "
        "suggested reorder point (lead-time demand + 95% safety stock).",
    ),
    "Technicians": (
        "SELECT * FROM v_mechanic_scorecard ORDER BY mechanic_id",
        "One row per technician: efficiency, on-time rate, comeback rate, CSAT, sales.",
    ),
    "MonthlyKPIs": (
        "SELECT * FROM v_monthly_kpis ORDER BY month",
        "One row per month: the executive KPI set.",
    ),
    "Customers": (
        "SELECT * FROM v_customer_summary ORDER BY customer_id",
        "One row per customer: acquisition channel, visits, lifetime sales, recency, satisfaction.",
    ),
}

DICTIONARY = [
    ("WorkOrders", "net_sales", "Labor + parts - discount, excluding sales tax"),
    ("WorkOrders", "gross_profit", "Net sales - parts cost - technician labor cost (hours_actual x hourly wage)"),
    ("WorkOrders", "is_on_time", "1 if completed_at <= promised_at; blank while the job is still open"),
    ("WorkOrders", "hours_late", "Hours past the promised time (0 when on time)"),
    ("WorkOrders", "turnaround_hours", "Clock hours from opened_at to completed_at"),
    ("WorkOrders", "hours_billed", "Flat-rate hours charged to the customer"),
    ("WorkOrders", "hours_actual", "Hours the technician actually clocked"),
    ("WorkOrders", "had_stockout", "1 if a normally stocked part was out and had to be special-ordered"),
    ("WorkOrders", "had_special_order", "1 if any part had to be ordered (stockout or never-stocked part)"),
    ("WorkOrders", "parts_wait_days", "Days between ordering and receiving the special-order part"),
    ("WorkOrders", "had_comeback", "1 if the car came back for a warranty rework of this job"),
    ("WorkOrders", "job_type", "Customer Pay, or Warranty Rework (no charge)"),
    ("WorkOrders", "checkin_wait_min", "Minutes between the appointment time and the job being opened"),
    ("WorkOrders", "csat_rating", "Post-visit survey, 1-5 stars (blank = no response)"),
    ("Appointments", "lead_days", "Days between booking and the appointment"),
    ("Appointments", "bookings_in_slot", "Appointments booked for the same start time"),
    ("PartsInventory", "on_hand", "Units received into stock minus units taken off the shelf"),
    ("PartsInventory", "suggested_reorder_point",
     "CEILING(avg daily use x avg actual lead days + 1.65 x std dev daily use x SQRT(lead days)), 2025 data"),
    ("PartsInventory", "times_out_of_stock", "Job lines where the part was not on the shelf"),
    ("Technicians", "efficiency_pct", "Hours billed / hours actual x 100 (over 100 = faster than book time)"),
    ("Technicians", "comeback_rate_pct", "Customer-pay jobs that needed a warranty rework / customer-pay jobs"),
    ("MonthlyKPIs", "avg_repair_order", "Net sales / customer-pay work orders"),
    ("MonthlyKPIs", "no_show_pct", "No-shows / (arrived + no-show + cancelled) appointments"),
    ("Customers", "days_since_last_visit", "Days from last visit to 2025-12-31"),
]


def connect():
    return pymysql.connect(
        host=os.environ.get("MYSQL_HOST", "localhost"),
        user=os.environ.get("MYSQL_USER", "root"),
        password=os.environ.get("MYSQL_PASSWORD", ""),
        database="service_management",
        charset="utf8mb4",
    )


def to_excel_value(v):
    if isinstance(v, decimal.Decimal):
        return float(v)
    return v


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["Service Management DB - dashboard dataset"])
    readme.append(["Synthetic data for an independent auto repair shop, 2024-01-01 to 2025-12-31 "
                   "(snapshot 2025-12-31 12:00). Generated by python/generate_data.py."])
    readme.append([])
    readme.append(["Sheet", "Rows", "Description"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    conn = connect()
    with conn.cursor() as cur:
        for sheet, (query, desc) in DATASETS.items():
            cur.execute(query)
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            with open(OUT / f"{sheet}.csv", "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(cols)
                w.writerows([["" if v is None else v for v in r] for r in rows])
            ws = wb.create_sheet(sheet)
            ws.append(cols)
            for r in rows:
                ws.append([to_excel_value(v) for v in r])
            for i, c in enumerate(cols, start=1):
                cell = ws.cell(row=1, column=i)
                cell.font, cell.fill = header_font, header_fill
                ws.column_dimensions[get_column_letter(i)].width = max(12, min(34, len(c) + 4))
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            readme.append([sheet, len(rows), desc])
            print(f"  {sheet:<16} {len(rows):>6,} rows")
    conn.close()

    readme.append([])
    readme.append(["Field definitions"])
    readme.append(["Sheet", "Field", "Definition"])
    for row in DICTIONARY:
        readme.append(list(row))
    for cell in (readme["A1"], readme["A4"], readme["B4"], readme["C4"]):
        cell.font = Font(bold=True)
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 26
    readme.column_dimensions["C"].width = 110
    wb.save(OUT / "ServiceOps_Dashboard_Data.xlsx")
    print(f"  wrote {OUT / 'ServiceOps_Dashboard_Data.xlsx'}")


if __name__ == "__main__":
    main()
